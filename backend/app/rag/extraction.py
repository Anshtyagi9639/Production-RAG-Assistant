"""Text extraction per file type. CPU-bound: call from a thread in async code."""

import io
import logging
import zipfile
from dataclasses import dataclass

import docx
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.core.errors import UserFacingError

logger = logging.getLogger(__name__)


class ExtractionError(UserFacingError):
    pass


@dataclass(frozen=True)
class Page:
    # 1-based page number for PDFs; None for formats without pages.
    number: int | None
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[Page]
    page_count: int | None


def _extract_pdf(data: bytes) -> ExtractedDocument:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("The PDF is password-protected.")
        pages = [
            Page(number, page.extract_text() or "")
            for number, page in enumerate(reader.pages, start=1)
        ]
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, OSError) as exc:
        logger.warning("pdf_extraction_failed", extra={"error": type(exc).__name__})
        raise ExtractionError("The PDF is corrupted or not a valid PDF file.") from exc
    return ExtractedDocument(pages, page_count=len(pages))


def _extract_docx(data: bytes) -> ExtractedDocument:
    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        logger.warning("docx_extraction_failed", extra={"error": type(exc).__name__})
        raise ExtractionError("The DOCX file is corrupted or not a valid Word document.") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return ExtractedDocument([Page(None, "\n".join(blocks))], page_count=None)


def _extract_text(data: bytes) -> ExtractedDocument:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ExtractionError("The file is not valid UTF-8 text.") from exc
    return ExtractedDocument([Page(None, text)], page_count=None)


def extract(data: bytes, extension: str) -> ExtractedDocument:
    if extension == ".pdf":
        return _extract_pdf(data)
    if extension == ".docx":
        return _extract_docx(data)
    if extension in {".txt", ".md"}:
        return _extract_text(data)
    raise ExtractionError(f"Unsupported file type: {extension}")
